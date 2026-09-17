"""Predict TestPicture images with a trained EfficientNet-B0 and render labeled PNG grids.
Run from any working directory. No training or weight updates are performed.
"""
import argparse
import csv
import hashlib
import json
import math
from pathlib import Path
import unicodedata

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont, ImageOps
import torch
from torch import nn
from torchvision import models, transforms

ROOT = Path(__file__).resolve().parents[1]
EXTENSIONS = {'.jpg', '.jpeg', '.png', '.webp', '.bmp', '.tif', '.tiff'}


def normalized(value):
    return unicodedata.normalize('NFC', str(value).strip().replace('\\', '/'))


def match_labels(images, labels):
    workbook = load_workbook(labels, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        columns = [str(x).strip() for x in next(iterator)]
        if len(set(columns)) != len(columns) or not {'filename', 'score'} <= set(columns):
            raise ValueError('Excel must have unique filename and score columns.')
        by_relative, by_name = {}, {}
        for path in images:
            relative = normalized('/'.join(path.parts[-2:]))
            by_relative[relative] = path
            by_name.setdefault(normalized(path.name), []).append(path)
        result, used, errors = [], set(), []
        for number, row in enumerate(iterator, 2):
            if all(value is None for value in row):
                continue
            try:
                name = row[columns.index('filename')]
                score = float(row[columns.index('score')])
                if name is None or not math.isfinite(score) or not 1 <= score <= 5:
                    raise ValueError('Missing filename or invalid score (expected 1..5).')
                name = normalized(name)
                candidates = [by_relative[name]] if name in by_relative else by_name.get(name, [])
                if len(candidates) != 1:
                    raise ValueError(f'Missing or ambiguous image: {name} ({len(candidates)} matches)')
                path = candidates[0]
                if path in used:
                    raise ValueError(f'Duplicate label for {name}')
                used.add(path)
                result.append({'path': path, 'filename': path.name, 'folder': path.parent.name, 'actual_score': score})
            except (ValueError, TypeError) as error:
                errors.append(f'Excel row {number}: {error}')
        for path in set(images) - used:
            errors.append(f'Image has no Excel label: {path.name}')
        if errors:
            raise ValueError('\n'.join(errors))
        return sorted(result, key=lambda row: (row['folder'], normalized(row['filename'])))
    finally:
        workbook.close()


def load_model(checkpoint, device):
    with torch.serialization.safe_globals([torch.torch_version.TorchVersion]):
        data = torch.load(checkpoint, map_location='cpu', weights_only=True)
    config = data.get('config', {})
    # Fail rather than silently change preprocessing/output for an unrelated checkpoint.
    if config.get('model') != 'efficientnet_b0' or config.get('task') != 'regression':
        raise ValueError('Expected an EfficientNet-B0 regression checkpoint with config metadata.')
    if config.get('output') != '1 + 4 * sigmoid':
        raise ValueError('Checkpoint output contract differs from 1 + 4 * sigmoid.')
    if config.get('preprocessing') != 'EXIF transpose, RGB, resize 256 bicubic, center crop 224':
        raise ValueError('Checkpoint preprocessing changed. Update this script to match training.')
    model = models.efficientnet_b0(weights=None)
    model.classifier[1] = nn.Linear(model.classifier[1].in_features, 1)
    model.load_state_dict(data['state_dict'], strict=True)
    model.to(device).eval()
    transform = transforms.Compose([
        transforms.Resize(256, interpolation=transforms.InterpolationMode.BICUBIC),
        transforms.CenterCrop(224), transforms.ToTensor(),
        transforms.Normalize(config['normalization_mean'], config['normalization_std']),
    ])
    return model, transform, config


def get_font(size, path=None):
    candidates = [path] if path else []
    candidates += ['C:/Windows/Fonts/malgun.ttf', '/usr/share/fonts/truetype/nanum/NanumGothic.ttf']
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return ImageFont.truetype(str(candidate), size)
    raise FileNotFoundError('Korean font not found. Pass --font path/to/KoreanFont.ttf')


def draw_grid(rows, target, title, columns, font_path):
    width, height, header = 320, 388, 94
    canvas = Image.new('RGB', (columns * width, header + math.ceil(len(rows) / columns) * height), '#f0f4fa')
    draw = ImageDraw.Draw(canvas)
    font, small, heading = get_font(21, font_path), get_font(14, font_path), get_font(27, font_path)
    draw.text((22, 12), title, font=heading, fill='#203450')
    draw.text((22, 54), 'EfficientNet-B0 | 실제 점수: 엑셀 라벨 | 예측 점수: 1~5', font=font, fill='#47658b')
    for i, row in enumerate(rows):
        x, y = (i % columns) * width + 10, header + (i // columns) * height
        draw.rounded_rectangle((x, y, x + 300, y + height - 12), radius=12, fill='white')
        with Image.open(row['path']) as image:
            photo = ImageOps.contain(ImageOps.exif_transpose(image).convert('RGB'), (284, 270))
            canvas.paste(photo, (x + (300-photo.width)//2, y + 8 + (270-photo.height)//2))
        label = row['filename']
        while draw.textbbox((0, 0), label, font=small)[2] > 278:
            label = label[:-2] + '…'
        draw.text((x + 12, y + 282), label, font=small, fill='#657994')
        draw.text((x + 12, y + 307), f"예측 {row['prediction']:.2f} / 실제 {row['actual_score']:.2f}", font=font, fill='#244b7b')
        draw.text((x + 12, y + 339), f"절대 오차 {row['absolute_error']:.2f}", font=small, fill='#657994')
    canvas.save(target)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--images', type=Path, default=ROOT / 'TestPicture')
    parser.add_argument('--labels', type=Path, default=ROOT / 'test_labels_1to5.xlsx')
    parser.add_argument('--checkpoint', type=Path, default=Path(__file__).resolve().parent / 'EfficientNetAllImage.pt')
    parser.add_argument('--output', type=Path, default=Path(__file__).resolve().parent / 'outputs' / 'test_picture')
    parser.add_argument('--columns', type=int, default=4)
    parser.add_argument('--batch-size', type=int, default=8)
    parser.add_argument('--font')
    parser.add_argument('--device', choices=['auto', 'cpu', 'cuda'], default='auto')
    args = parser.parse_args()
    if not 1 <= args.columns <= 8 or args.batch_size < 1:
        parser.error('columns must be 1..8; batch-size must be positive')
    get_font(20, args.font)
    images = []
    for group in range(1, 6):
        folder = args.images / str(group)
        group_images = sorted(p for p in folder.rglob('*') if p.is_file() and p.suffix.lower() in EXTENSIONS)
        if not group_images:
            raise ValueError(f'No images in {folder}')
        images.extend(group_images)
    rows = match_labels(images, args.labels)
    device = torch.device(('cuda' if torch.cuda.is_available() else 'cpu') if args.device == 'auto' else args.device)
    torch.set_num_threads(min(4, torch.get_num_threads()))
    model, transform, config = load_model(args.checkpoint, device)
    print(f'Matched {len(rows)} images and Excel labels; device={device}', flush=True)
    with torch.inference_mode():
        for start in range(0, len(rows), args.batch_size):
            batch = rows[start:start + args.batch_size]
            tensors = []
            for row in batch:
                with Image.open(row['path']) as image:
                    tensors.append(transform(ImageOps.exif_transpose(image).convert('RGB')))
            predictions = (1 + 4 * model(torch.stack(tensors).to(device)).sigmoid()).flatten().cpu().tolist()
            for row, prediction in zip(batch, predictions):
                if not math.isfinite(prediction):
                    raise ValueError(f"Nonfinite prediction: {row['path']}")
                row['prediction'] = prediction
                row['absolute_error'] = abs(prediction - row['actual_score'])
            print(f'Predicted {min(start + args.batch_size, len(rows))}/{len(rows)}', flush=True)
    args.output.mkdir(parents=True, exist_ok=True)
    for group in range(1, 6):
        group_rows = [r for r in rows if r['folder'] == str(group)]
        for start in range(0, len(group_rows), 12):
            page = start // 12 + 1
            draw_grid(group_rows[start:start+12], args.output / f'score_{group}_{page:02d}.png',
                      f'TestPicture / {group} · {page} 페이지', args.columns, args.font)
    fields = ['folder', 'filename', 'actual_score', 'prediction', 'absolute_error']
    with (args.output / 'predictions.csv').open('w', encoding='utf-8-sig', newline='') as file:
        writer = csv.DictWriter(file, fieldnames=fields, extrasaction='ignore')
        writer.writeheader(); writer.writerows(rows)
    mse = sum(r['absolute_error'] ** 2 for r in rows) / len(rows)
    actual = torch.tensor([r['actual_score'] for r in rows], dtype=torch.float64)
    pred = torch.tensor([r['prediction'] for r in rows], dtype=torch.float64)
    pearson = torch.corrcoef(torch.stack([actual, pred]))[0, 1].item()
    metrics = {'n': len(rows), 'mse': mse, 'mae': sum(r['absolute_error'] for r in rows)/len(rows),
               'rmse': math.sqrt(mse), 'pearson_r': pearson if math.isfinite(pearson) else None,
               'checkpoint': str(args.checkpoint.resolve()),
               'checkpoint_sha256': hashlib.sha256(args.checkpoint.read_bytes()).hexdigest(),
               'labels': str(args.labels.resolve()), 'training_config': config}
    (args.output / 'metrics.json').write_text(json.dumps(metrics, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f"Test n={len(rows)} MSE={mse:.4f} MAE={metrics['mae']:.4f} RMSE={metrics['rmse']:.4f} Pearson r={pearson:.4f}")
    print(f'PNG / CSV / metrics: {args.output.resolve()}')

if __name__ == '__main__':
    main()
