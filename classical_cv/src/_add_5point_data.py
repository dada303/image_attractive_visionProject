"""일회성 스크립트: custom_dataset/5_dataset의 5점 이미지를 images/5_male, images/5_female로
복사하고, labels_combined_1to5.xlsx에 5점 라벨로 추가한다.

용도가 끝나면 삭제해도 되는 1회성 유틸리티 (파이프라인 코드가 아님).
"""
import os
import shutil
import openpyxl

BASE = r"C:\Users\한국전파진흥협회\Desktop\dataset"
SRC = os.path.join(BASE, "custom_dataset", "5_dataset")
IMAGES = os.path.join(BASE, "combined", "images")
XLSX = os.path.join(BASE, "combined", "labels_combined_1to5.xlsx")

# (원본 하위 폴더, 목적지 카테고리, 파일명 접두어 또는 None(원본 유지))
SOURCES = [
    (os.path.join(SRC, "C-MALE 5start", "C-MALE 5start"), "5_male", "cn_m_"),
    (os.path.join(SRC, "K-MALE 5star"), "5_male", "kr_m_"),
    (os.path.join(SRC, "cropped_faces_japen", "cropped_faces_japen"), "5_male", None),
    (os.path.join(SRC, "여자데이터셋 5점", "여자데이터셋 5점"), "5_female", "kr_f_"),
    (os.path.join(SRC, "chinese_5_FM", "chinese_5_FM"), "5_female", "cn_f_"),
]

# 확인된 정확한 중복 파일(md5 동일) - 복사에서 제외
SKIP_BASENAMES = {"3227d0b5ed315b7d255516f18bff26ba (1).jpg"}


def sanitize(name):
    return name.replace(" ", "_").replace("(", "").replace(")", "")


def copy_all():
    added = {"5_male": [], "5_female": []}
    for src_dir, category, prefix in SOURCES:
        for fname in sorted(os.listdir(src_dir)):
            src_path = os.path.join(src_dir, fname)
            if not os.path.isfile(src_path):
                continue
            if fname in SKIP_BASENAMES:
                print(f"[skip-dup] {fname}")
                continue
            dst_name = sanitize((prefix + fname) if prefix else fname)
            dst_path = os.path.join(IMAGES, category, dst_name)
            shutil.copyfile(src_path, dst_path)
            added[category].append(dst_name)
    return added


def verify_decodable(added):
    import numpy as np
    import cv2

    fail = []
    for category, names in added.items():
        for name in names:
            path = os.path.join(IMAGES, category, name)
            data = np.fromfile(path, dtype=np.uint8)
            img = cv2.imdecode(data, cv2.IMREAD_COLOR)
            if img is None:
                fail.append(path)
    return fail


def update_excel(added):
    wb = openpyxl.load_workbook(XLSX)
    total_sheet = wb["전체"]

    for category in ("5_male", "5_female"):
        if category in wb.sheetnames:
            raise RuntimeError(
                f"'{category}' 시트가 이미 존재합니다 — 이 스크립트를 이미 실행한 것 같습니다. "
                "중복 추가를 막기 위해 중단합니다."
            )

    for category in ("5_male", "5_female"):
        names = added[category]
        ws = wb.create_sheet(category)
        ws.append(["filename", "score"])
        for name in names:
            ws.append([name, 5])

        for name in names:
            total_sheet.append([category, name, 5, "1~5"])

    wb.save(XLSX)


if __name__ == "__main__":
    added = copy_all()
    print(f"5_male: {len(added['5_male'])}개, 5_female: {len(added['5_female'])}개")

    fail = verify_decodable(added)
    if fail:
        print(f"[경고] 디코딩 실패 {len(fail)}건:")
        for p in fail:
            print(" ", p)
    else:
        print("모든 복사 파일 cv2 디코딩 확인 완료.")

    update_excel(added)
    print("labels_combined_1to5.xlsx 업데이트 완료.")
