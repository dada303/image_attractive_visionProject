"""일회성 스크립트: idolgirlcropsimage/ 의 이미지(+ 동봉된 labels.xlsx 라벨)를
images/5_female로 복사하고, labels_combined_1to5.xlsx의 5_female 시트/전체 시트에 추가한다.

파이프라인 코드가 아닌 1회성 유틸리티.
"""
import os
import shutil
import openpyxl

BASE = r"C:\Users\한국전파진흥협회\Desktop\dataset"
SRC_DIR = os.path.join(BASE, "idolgirlcropsimage")
SRC_LABELS = os.path.join(SRC_DIR, "labels.xlsx")
IMAGES = os.path.join(BASE, "combined", "images")
XLSX = os.path.join(BASE, "combined", "labels_combined_1to5.xlsx")

CATEGORY = "5_female"
PREFIX = "idol_"


def sanitize(name):
    return name.replace(" ", "_").replace("(", "").replace(")", "")


def read_source_labels():
    wb = openpyxl.load_workbook(SRC_LABELS, data_only=True)
    ws = wb["labels"]
    return [(row[0], row[1]) for row in ws.iter_rows(min_row=2, values_only=True)]


def copy_all(labels):
    added = []
    for fname, score in labels:
        src_path = os.path.join(SRC_DIR, fname)
        if not os.path.isfile(src_path):
            raise FileNotFoundError(f"labels.xlsx에는 있는데 파일이 없음: {fname}")
        dst_name = sanitize(PREFIX + fname)
        dst_path = os.path.join(IMAGES, CATEGORY, dst_name)
        if os.path.exists(dst_path):
            raise FileExistsError(f"이미 존재하는 파일명: {dst_path}")
        shutil.copyfile(src_path, dst_path)
        added.append((dst_name, score))
    return added


def verify_decodable(added):
    import numpy as np
    import cv2

    fail = []
    for name, _ in added:
        path = os.path.join(IMAGES, CATEGORY, name)
        data = np.fromfile(path, dtype=np.uint8)
        img = cv2.imdecode(data, cv2.IMREAD_COLOR)
        if img is None:
            fail.append(path)
    return fail


def update_excel(added):
    wb = openpyxl.load_workbook(XLSX)
    total_sheet = wb["전체"]
    female_sheet = wb[CATEGORY]

    existing = {row[0] for row in female_sheet.iter_rows(min_row=2, values_only=True)}
    dup = [name for name, _ in added if name in existing]
    if dup:
        raise RuntimeError(f"'{CATEGORY}' 시트에 이미 있는 파일명이 섞여 있습니다: {dup[:5]}...")

    for name, score in added:
        female_sheet.append([name, score])
        total_sheet.append([CATEGORY, name, score, "1~5"])

    wb.save(XLSX)


if __name__ == "__main__":
    labels = read_source_labels()
    print(f"소스 labels.xlsx: {len(labels)}건")

    added = copy_all(labels)
    print(f"{CATEGORY}에 복사 완료: {len(added)}장")

    fail = verify_decodable(added)
    if fail:
        print(f"[경고] 디코딩 실패 {len(fail)}건:")
        for p in fail:
            print(" ", p)
    else:
        print("모든 복사 파일 cv2 디코딩 확인 완료.")

    update_excel(added)
    print("labels_combined_1to5.xlsx 업데이트 완료.")
